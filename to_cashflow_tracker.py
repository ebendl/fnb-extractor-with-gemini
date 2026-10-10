#!/usr/bin/env python3
"""
Cashflow Tracker Import Converter

Converts the CSV/Excel files produced by this repo into the Excel-format CSV
that cashflow_tracker's import endpoints (and seed runner) expect.

Input  (either of the two formats this repo produces):
  - extractor CSV : date;description;amount;balance          (semicolon separated)
  - converted xlsx: Date, Debit amount, Credit amount, Saldo,
                    Calculated saldo, transaction description

Output columns (cashflow_tracker Excel-format CSV):
  Salary Month, Salary, Date, Debit, Credit, Saldo, Calc Saldo,
  Bank Note, Split, Category, Booked, Note

'Salary Month' is deliberately left blank — it is populated by hand. Note that
cashflow_tracker's parser SKIPS rows with a blank Salary Month, and infers the
transaction year from it, so the column must be filled in before importing.
"""

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

# Header of the cashflow_tracker Excel-format CSV, in order.
OUTPUT_COLUMNS = [
    "Salary Month", "Salary", "Date", "Debit", "Credit", "Saldo",
    "Calc Saldo", "Bank Note", "Split", "Category", "Booked", "Note",
]

# Written in the Debit/Credit column when that side of the pair is absent.
# cashflow_tracker's parse_amount() reads a bare dash as "zero / not applicable".
EMPTY_AMOUNT = "-"

OUTPUT_SUFFIX = "_cashflow.csv"

# Produced by combine_csv_files.py: it duplicates the individual statement
# rows, so it must never be folded into a combined conversion.
COMBINED_STEM = "combined_transactions"

EXTRACTOR_CSV_COLUMNS = ["date", "description", "amount", "balance"]
EXCEL_COLUMNS = [
    "Date", "Debit amount", "Credit amount", "Saldo",
    "Calculated saldo", "transaction description",
]


@dataclass
class Row:
    """One statement line, normalised across both input formats."""

    txn_date: date
    description: str
    amount: Decimal   # signed: negative = debit, positive = credit
    balance: Decimal


def to_decimal(value, field: str, source: str) -> Decimal:
    """Coerce a CSV string or Excel float to a 2dp Decimal."""
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"{source}: empty '{field}'")
    try:
        # str() via float keeps openpyxl's binary floats from leaking artefacts.
        return Decimal(str(value).strip()).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{source}: cannot read '{field}' value {value!r}")


def to_date(value, source: str) -> date:
    """Accept an ISO date string or an Excel/pandas datetime."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        raise ValueError(f"{source}: empty date")
    # Excel round-trips can append a time component.
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        raise ValueError(f"{source}: cannot read date {value!r}")


def read_extractor_csv(path: Path) -> list[Row]:
    """Read the semicolon-separated CSV written by bank_statement_extractor.py."""
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, delimiter=";")
        fields = [f.strip() for f in (reader.fieldnames or [])]
        missing = [c for c in EXTRACTOR_CSV_COLUMNS if c not in fields]
        if missing:
            raise ValueError(
                f"{path.name}: not an extractor CSV (missing {', '.join(missing)})"
            )

        rows = []
        for line_no, raw in enumerate(reader, start=2):
            record = {(k.strip() if k else k): v for k, v in raw.items()}
            if not (record.get("date") or "").strip():
                continue  # blank trailing line
            source = f"{path.name} line {line_no}"
            rows.append(
                Row(
                    txn_date=to_date(record["date"], source),
                    description=(record.get("description") or "").strip(),
                    amount=to_decimal(record["amount"], "amount", source),
                    balance=to_decimal(record["balance"], "balance", source),
                )
            )
    return rows


def read_converted_excel(path: Path) -> list[Row]:
    """Read the .xlsx written by csv_to_excel.py."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        grid = sheet.iter_rows(values_only=True)
        try:
            header = [str(c).strip() if c is not None else "" for c in next(grid)]
        except StopIteration:
            raise ValueError(f"{path.name}: empty worksheet")

        missing = [c for c in EXCEL_COLUMNS if c not in header]
        if missing:
            raise ValueError(
                f"{path.name}: not a converted statement workbook "
                f"(missing {', '.join(missing)})"
            )
        index = {name: header.index(name) for name in EXCEL_COLUMNS}

        rows = []
        for row_no, cells in enumerate(grid, start=2):
            if cells[index["Date"]] in (None, ""):
                continue  # blank trailing row
            source = f"{path.name} row {row_no}"
            debit = cells[index["Debit amount"]] or 0
            credit = cells[index["Credit amount"]] or 0
            # csv_to_excel.py splits the signed amount into two columns, writing
            # 0.0 in the unused one; the debit column keeps its negative sign.
            amount = (
                to_decimal(debit, "Debit amount", source)
                + to_decimal(credit, "Credit amount", source)
            )
            rows.append(
                Row(
                    txn_date=to_date(cells[index["Date"]], source),
                    description=str(cells[index["transaction description"]] or "").strip(),
                    amount=amount,
                    balance=to_decimal(cells[index["Saldo"]], "Saldo", source),
                )
            )
    finally:
        workbook.close()
    return rows


def read_statement(path: Path) -> list[Row]:
    """Dispatch on file extension."""
    if path.suffix.lower() == ".csv":
        return read_extractor_csv(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return read_converted_excel(path)
    raise ValueError(f"{path.name}: unsupported file type '{path.suffix}'")


def to_output_row(row: Row) -> dict:
    """Map a normalised row onto the cashflow_tracker columns.

    A row is either a debit or a credit; the unused side gets a dash. A zero
    amount (fee notices, reversals) leaves both sides empty, matching how the
    historical exports record them.
    """
    debit = credit = EMPTY_AMOUNT
    if row.amount < 0:
        debit = f"{row.amount:.2f}"    # kept negative, as in the historical exports
    elif row.amount > 0:
        credit = f"{row.amount:.2f}"

    balance = f"{row.balance:.2f}"
    return {
        "Salary Month": "",           # populated by hand
        "Salary": "",
        "Date": f"{row.txn_date.day:02d} {row.txn_date.strftime('%b')}",
        "Debit": debit,
        "Credit": credit,
        "Saldo": balance,
        "Calc Saldo": balance,
        "Bank Note": row.description,
        "Split": "",
        "Category": "",
        "Booked": "",
        "Note": "",
    }


def write_output(rows: list[Row], output_path: Path) -> None:
    """Write the Excel-format CSV.

    utf-8-sig so Excel opens it cleanly; the import endpoint decodes with
    utf-8-sig too, so the BOM is stripped on the way in.
    """
    with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow(to_output_row(row))


def natural_key(path: Path):
    """Sort FIRST_..._58 before FIRST_..._100 (statement numbers, not strings)."""
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", path.stem)
    ]


def find_inputs(directory: Path) -> list[Path]:
    """Collect statement files in a directory, one per statement.

    Skips this script's own output and the combine_csv_files.py output, and
    prefers the CSV when a statement has both a .csv and an .xlsx.
    """
    by_stem: dict[str, Path] = {}
    for path in sorted(directory.iterdir(), key=natural_key):
        if not path.is_file() or path.suffix.lower() not in (".csv", ".xlsx", ".xlsm"):
            continue
        if path.name.endswith(OUTPUT_SUFFIX) or path.stem == COMBINED_STEM:
            continue
        if path.name.startswith("~$"):
            continue  # Excel lock file
        # CSV wins over the .xlsx generated from it.
        if path.stem not in by_stem or path.suffix.lower() == ".csv":
            by_stem[path.stem] = path
    return sorted(by_stem.values(), key=natural_key)


def convert_directory(directory: Path, combined: bool, output: Path | None) -> bool:
    inputs = find_inputs(directory)
    if not inputs:
        print(f"No statement CSV/Excel files found in: {directory}")
        return False

    print(f"🔍 Found {len(inputs)} statement files in {directory}")

    if not combined:
        converted = 0
        for path in inputs:
            if convert_file(path, path.with_name(path.stem + OUTPUT_SUFFIX)):
                converted += 1
        print(f"\n✨ Converted {converted} of {len(inputs)} files.")
        return converted == len(inputs)

    all_rows: list[Row] = []
    for path in inputs:
        try:
            rows = read_statement(path)
        except (ValueError, OSError) as error:
            print(f"❌ {error}")
            return False
        print(f"   📄 {path.name}: {len(rows)} transactions")
        all_rows.extend(rows)

    # Statement files are already in statement order, and their numbering runs
    # chronologically, so concatenating in natural order preserves the running
    # balance. Report overlaps rather than dropping them — cashflow_tracker's
    # import dedups, and a silent drop would hide a genuinely repeated fee.
    seen: set[tuple] = set()
    duplicates = 0
    for row in all_rows:
        identity = (row.txn_date, row.description, row.amount, row.balance)
        if identity in seen:
            duplicates += 1
        seen.add(identity)
    if duplicates:
        print(f"   ⚠️  {duplicates} identical rows across statements (overlapping periods?)")

    target = output or directory / (
        re.sub(r"\s+", "_", directory.name.strip()) + OUTPUT_SUFFIX
    )
    write_output(all_rows, target)
    print(f"\n✅ Wrote {len(all_rows)} transactions to: {target}")
    print("   ⚠️  'Salary Month' is blank — fill it in before importing.")
    return True


def convert_file(input_path: Path, output_path: Path) -> bool:
    try:
        rows = read_statement(input_path)
    except (ValueError, OSError) as error:
        print(f"❌ {error}")
        return False

    write_output(rows, output_path)
    print(f"✅ {input_path.name}: {len(rows)} transactions → {output_path.name}")
    return True


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert this repo's statement CSV/Excel files into the "
                    "Excel-format CSV that cashflow_tracker imports.",
    )
    parser.add_argument(
        "input",
        help="Statement .csv/.xlsx file, or a directory of them (one account).",
    )
    parser.add_argument(
        "-o", "--output",
        help="Output CSV path. Defaults to <name>%s next to the input." % OUTPUT_SUFFIX,
    )
    parser.add_argument(
        "--per-file",
        action="store_true",
        help="For a directory, write one output per statement instead of a "
             "single combined file.",
    )
    args = parser.parse_args()

    input_path = Path(args.input).expanduser()
    output = Path(args.output).expanduser() if args.output else None

    if input_path.is_dir():
        if output and args.per_file:
            parser.error("--output cannot be combined with --per-file")
        ok = convert_directory(input_path, combined=not args.per_file, output=output)
    elif input_path.is_file():
        target = output or input_path.with_name(input_path.stem + OUTPUT_SUFFIX)
        ok = convert_file(input_path, target)
        if ok:
            print("   ⚠️  'Salary Month' is blank — fill it in before importing.")
    else:
        print(f"Error: path not found: {input_path}")
        sys.exit(1)

    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
